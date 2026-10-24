import json, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

PALETTES = {
 'academic': ('142B49','117C83','F4F8FA'),
 'business': ('14233C','2865C7','F5F7FB'),
 'technical': ('17232E','087F9C','F1F6F8'),
 'creative': ('40244D','D15E6E','FBF5F7')}

def rgb(hexstr): return RGBColor.from_string(hexstr)
def rect(slide,x,y,w,h,color,rounded=False):
 s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 s.fill.solid(); s.fill.fore_color.rgb=rgb(color); s.line.fill.background(); return s

def txt(slide,text,x,y,w,h,size=20,color='17232E',bold=False,align=None):
 shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=shape.text_frame; tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.06); tf.margin_top=tf.margin_bottom=Inches(.03)
 for i,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
  if align is not None:p.alignment=align
 return shape

def items(slide,arr,x,y,w,size=19,step=.7):
 for i,s in enumerate(arr[:7]):
  txt(slide,'• '+str(s),x,y+i*step,w,step-.05,size)

def main():
 if len(sys.argv)!=3: sys.exit('Usage: python generate_ppt.py deck.json output.pptx')
 spec=json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
 prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
 dark,accent,bg=PALETTES.get(spec.get('theme','business'),PALETTES['business'])
 custom=spec.get('palette',{})
 if custom:
  for k in ('dark','accent','background'):
   v=custom.get(k)
   if v is not None and (not isinstance(v,str) or len(v.lstrip('#'))!=6 or any(c not in '0123456789abcdefABCDEF' for c in v.lstrip('#'))):
    raise ValueError('Invalid hex palette color: '+k)
  dark=custom.get('dark',dark).lstrip('#'); accent=custom.get('accent',accent).lstrip('#'); bg=custom.get('background',bg).lstrip('#')
 def luminance(h):
  channels=[int(h[i:i+2],16)/255 for i in (0,2,4)]
  channels=[c/12.92 if c<=.04045 else ((c+.055)/1.055)**2.4 for c in channels]
  return sum(a*b for a,b in zip(channels,(.2126,.7152,.0722)))
 contrast=(max(luminance(dark),luminance(bg))+.05)/(min(luminance(dark),luminance(bg))+.05)
 if contrast<4.5: raise ValueError(f'Insufficient text/background contrast: {contrast:.2f}:1')
 for idx,d in enumerate(spec['slides']):
  slide=prs.slides.add_slide(prs.slide_layouts[6]); rect(slide,0,0,13.333,7.5,bg)
  layout=d['layout']; title=d.get('title','')
  if layout in ('title','closing'):
   rect(slide,0,0,.23,7.5,accent); txt(slide,title,.9,2.15,11.5,1.25,42,dark,True)
   txt(slide,d.get('subtitle',''),.95,3.6,11.2,1.4,24,accent)
  else:
   rect(slide,.65,.53,.12,.6,accent); txt(slide,title,.95,.48,11.6,.88,32,dark,True)
   rect(slide,.85,1.48,11.7,.025,accent)
   if layout=='bullets':items(slide,d.get('bullets',[]),1.05,1.9,11,21,.78)
   elif layout in ('two_column','comparison'):
    for x,key,head in [(1,'left','left_title'),(6.8,'right','right_title')]:
     rect(slide,x,1.9,5.45,4.75,'FFFFFF',True); txt(slide,d.get(head,''),x+.25,2.1,4.9,.55,23,accent,True)
     items(slide,d.get(key,[]),x+.25,2.95,4.85,18,.65)
   elif layout=='process':
    steps=d.get('steps',[])[:5]; n=max(1,len(steps)); width=11.4/n
    for i,step in enumerate(steps):
     x=.95+i*width; rect(slide,x,2.6,width-.18,2.05,'FFFFFF',True)
     txt(slide,str(i+1),x+.17,2.78,.6,.6,28,accent,True)
     txt(slide,step,x+.18,3.43,width-.52,1.0,17,dark)
   elif layout=='quote':
    txt(slide,'“'+d.get('quote','')+'”',1.2,2.15,10.9,2.7,29,dark)
    txt(slide,d.get('attribution',''),1.25,5.25,10,.6,18,accent)
   elif layout=='metrics':
    metrics=d.get('metrics',[])[:4]; n=max(1,len(metrics)); width=11.5/n
    for i,m in enumerate(metrics):
     x=.9+i*width; rect(slide,x,2.45,width-.2,2.7,'FFFFFF',True)
     txt(slide,m.get('value',''),x+.12,2.95,width-.45,.85,34,accent,True,PP_ALIGN.CENTER)
     txt(slide,m.get('label',''),x+.12,4.05,width-.45,.8,17,dark,False,PP_ALIGN.CENTER)
   elif layout=='chart':
    cats=d.get('categories',[]); vals=d.get('values',[])
    if not cats or len(cats)!=len(vals): raise ValueError('Chart categories/values mismatch')
    cd=CategoryChartData(); cd.categories=[str(c) for c in cats]; cd.add_series(d.get('chart_title','Series'),[float(v) for v in vals])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.25), Inches(1.9), Inches(10.75), Inches(4.7),cd).chart
    chart.has_legend=False; chart.has_title=bool(d.get('chart_title')); chart.chart_title.text_frame.text=d.get('chart_title','') if chart.has_title else ''
    txt(slide,d.get('unit',''),10.5,6.55,1.6,.4,13,accent)
   else:raise ValueError('Unsupported layout: '+layout)
   txt(slide,f'{idx+1:02d}',12.25,7.02,.5,.3,12,accent)
  if d.get('speaker_notes'):
   slide.notes_slide.notes_text_frame.text=d['speaker_notes']
 out=Path(sys.argv[2]); out.parent.mkdir(parents=True,exist_ok=True); prs.save(str(out)); print('Created',out)
if __name__=='__main__':main()
