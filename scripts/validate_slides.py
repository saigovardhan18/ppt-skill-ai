import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def main():
 if len(sys.argv)!=2:sys.exit('Usage: python validate_slides.py output.pptx')
 p=Presentation(sys.argv[1]); issues=[]
 if not p.slides:issues.append('No slides')
 for i,s in enumerate(p.slides,1):
  texts=[]
  for sh in s.shapes:
   if sh.left < -1000 or sh.top < -1000 or sh.left+sh.width > p.slide_width+1000 or sh.top+sh.height > p.slide_height+1000:issues.append(f'Slide {i}: shape outside bounds')
   if sh.has_text_frame:
    t=sh.text.strip(); texts.append(t)
    if len(t)>450:issues.append(f'Slide {i}: dense text block ({len(t)} characters)')
  if not any(texts):issues.append(f'Slide {i}: no text')
 print(f'Validated {len(p.slides)} slides; {len(issues)} warning(s)')
 for issue in issues:print('WARNING:',issue)
 if issues:sys.exit(1)
if __name__=='__main__':main()
